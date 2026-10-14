"""Extract supplied workbooks without modifying them.

Usage: python scripts/import_sources.py --manifest .local/source_manifest.json
Requires openpyxl only for XLSX reading. Public records omit collector names,
examples, comments and absolute source paths. Full cell snapshots stay local.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PRIVATE_FIELDS = {"收集人", "示例"}


def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def value(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v


def extract(manifest):
    sources, records, snapshots = [], [], []
    for src in manifest["sources"]:
        path = Path(src["path"])
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        wb = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
        sheets, source_rows = [], 0
        for ws in wb:
            raw_rows = []
            header_row, headers = None, {}
            for row in ws:
                cells = []
                for c in row:
                    if c.value is None and not c.hyperlink:
                        continue
                    cell = {"cell": c.coordinate, "value": value(c.value)}
                    if c.hyperlink:
                        cell["hyperlink"] = c.hyperlink.target or c.hyperlink.location
                    if c.data_type == "f":
                        cell["cached_value"] = value(cached[ws.title][c.coordinate].value)
                    cells.append(cell)
                if cells:
                    raw_rows.append({"row": row[0].row, "cells": cells})
                if header_row is None and any(c.value == "数据集名称" for c in row):
                    header_row = row[0].row
                    headers = {c.column: str(c.value).strip() for c in row if c.value is not None}
            sheet_rows = 0
            if header_row:
                name_col = next(k for k, v in headers.items() if v == "数据集名称")
                for row_num in range(header_row + 1, ws.max_row + 1):
                    name = ws.cell(row_num, name_col).value
                    if name is None or not str(name).strip():
                        continue
                    fields, links, coordinates = {}, [], {}
                    for col, label in headers.items():
                        if label in PRIVATE_FIELDS:
                            continue
                        c = ws.cell(row_num, col)
                        v = value(c.value)
                        if c.data_type == "f":
                            v = value(cached[ws.title][c.coordinate].value)
                        if v is not None:
                            fields[label] = v
                            coordinates[label] = c.coordinate
                        if c.hyperlink and c.hyperlink.target:
                            links.append({"cell": c.coordinate, "url": c.hyperlink.target})
                    records.append({"record_id": f'{src["id"]}:{ws.title}:{row_num}',
                                    "source_id": src["id"], "sheet": ws.title,
                                    "row": row_num, "name": str(name).strip(),
                                    "fields": fields, "cells": coordinates, "hyperlinks": links})
                    sheet_rows += 1
            source_rows += sheet_rows
            sheets.append({"name": ws.title, "max_row": ws.max_row,
                           "max_column": ws.max_column, "header_row": header_row,
                           "named_rows": sheet_rows,
                           "nonempty_rows": len(raw_rows), "embedded_images": len(ws._images)})
            snapshots.append({"source_id": src["id"], "sheet": ws.title, "rows": raw_rows})
        wb.close()
        cached.close()
        after = hashlib.sha256(path.read_bytes()).hexdigest()
        if before != after:
            raise RuntimeError(f"Source changed during extraction: {path.name}")
        sources.append({"id": src["id"], "filename": path.name, "sha256": before,
                        "bytes": path.stat().st_size, "named_rows": source_rows, "sheets": sheets})
    dump(ROOT / "data/sources.json", {"imported_on": manifest["imported_on"], "sources": sources})
    dump(ROOT / "data/source_records.json", {"schema_version": "1.0", "records": records})
    dump(ROOT / ".local/source_snapshot.json", {"snapshots": snapshots})
    print(json.dumps({"sources": len(sources), "named_rows_including_repeated_views": len(records),
                      "sheets": sum(len(s["sheets"]) for s in sources)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    extract(json.loads(args.manifest.read_text(encoding="utf-8-sig")))
